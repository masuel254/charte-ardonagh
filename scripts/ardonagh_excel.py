#!/usr/bin/env python3
"""Charte Ardonagh pour tout classeur Excel (inspirée du Q&A Dune et du one-pager). Valeurs et formules jamais modifiées.
Usage :
  python3 ardonagh_excel.py inspect in.xlsx
  python3 ardonagh_excel.py apply in.xlsx out.xlsx [params.json]
  python3 ardonagh_excel.py check in.xlsx out.xlsx
"""
import re, os, sys, json, zipfile, shutil, tempfile, subprocess, colorsys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from ardonagh_word import classify, GREEN, DEEP, MID, SOFT, SOFT2, LINE, SUB, GOLD, RED, REDSOFT, BAR, THEME

RISKY = ('xl/charts/', 'xl/drawings/', 'xl/pivotTables/', 'xl/pivotCache/', 'xl/media/', 'xl/slicers/', 'xl/timelines/', 'xl/vbaProject.bin', 'xl/chartsheets/')

def risky_parts(path):
    return sorted({p for n in zipfile.ZipFile(path).namelist() for p in RISKY if n.startswith(p)})

# ------------------------------------------------------------------ mode 1 : palette (XML, sans perte)
def argb_map(m, kind):
    v = m.group(2).upper()
    if len(v) == 8: a, h = v[:2], v[2:]
    else: a, h = 'FF', v
    return f'{m.group(1)}{a}{classify(h, kind)}"'

def palette_xml(src, out):
    """Polices en Arial et couleurs ramenées dans la palette, directement dans le XML : graphiques, TCD et macros intacts."""
    wd = tempfile.mkdtemp(prefix='ardx_')
    with zipfile.ZipFile(src) as z:
        for n in z.namelist():
            if '..' in n or n.startswith('/'): continue
            z.extract(n, wd)
    p = lambda *a: os.path.join(wd, *a)
    st = open(p('xl', 'styles.xml'), encoding='utf-8').read()
    # polices
    st = re.sub(r'<name val="[^"]*"/>', '<name val="Arial"/>', st)
    st = re.sub(r'<scheme val="(?:minor|major)"/>', '', st)
    # couleurs de police / de remplissage / de bordure
    def blk(tag, kind, s):
        return re.sub(rf'(<{tag}\b.*?</{tag}>)', lambda m: re.sub(r'(rgb=")([0-9A-Fa-f]{6,8})"', lambda x: argb_map(x, kind), m.group(1)), s, flags=re.S)
    st = blk('font', 'text', st); st = blk('fill', 'fill', st); st = blk('border', 'border', st)
    st = re.sub(r'(<dxf>.*?</dxf>)', lambda m: blk('fill', 'fill', blk('font', 'text', m.group(1))), st, flags=re.S)
    open(p('xl', 'styles.xml'), 'w', encoding='utf-8').write(st)
    th = p('xl', 'theme', 'theme1.xml')
    if os.path.exists(th):
        t = open(th, encoding='utf-8').read()
        for k, v in THEME.items(): t = re.sub(rf'(<a:{k}>)(.*?)(</a:{k}>)', rf'\1<a:srgbClr val="{v}"/>\3', t, flags=re.S)
        t = re.sub(r'(<a:(?:major|minor)Font><a:latin typeface=")[^"]*"', r'\1Arial"', t)
        open(th, 'w', encoding='utf-8').write(t)
    # graphiques : les séries en couleurs de thème suivent le nouveau thème ; les couleurs explicites sont ramenées dans la palette
    if os.path.isdir(p('xl', 'charts')):
        for f in os.listdir(p('xl', 'charts')):
            if not f.endswith('.xml'): continue
            c = open(p('xl', 'charts', f), encoding='utf-8').read()
            c = re.sub(r'(<a:srgbClr val=")([0-9A-Fa-f]{6})"', lambda m: m.group(1) + classify(m.group(2), 'fill') + '"', c)
            c = re.sub(r'typeface="(?!\+)[^"]*"', 'typeface="Arial"', c)
            open(p('xl', 'charts', f), 'w', encoding='utf-8').write(c)
    for d, _, fs in os.walk(p('xl', 'worksheets')):
        for f in fs:
            if f.endswith('.xml'):
                x = open(os.path.join(d, f), encoding='utf-8').read()
                tab = f'<tabColor rgb="FF{GREEN}"/>'
                if '<tabColor' in x: x = re.sub(r'<tabColor [^>]*/>', tab, x)
                elif re.search(r'<sheetPr[^>]*/>', x): x = re.sub(r'<sheetPr([^>]*)/>', rf'<sheetPr\1>{tab}</sheetPr>', x, 1)
                elif '<sheetPr' in x: x = re.sub(r'(<sheetPr[^>]*>)', rf'\1{tab}', x, 1)
                else: x = re.sub(r'(<worksheet[^>]*>)', rf'\1<sheetPr>{tab}</sheetPr>', x, 1)
                x = re.sub(r'(<sheetView\b[^>]*?)\s+showGridLines="[^"]*"', r'\1', x)
                x = re.sub(r'<sheetView\b', '<sheetView showGridLines="0"', x, 1)
                open(os.path.join(d, f), 'w', encoding='utf-8').write(x)
    if os.path.exists(out): os.remove(out)
    subprocess.run(f'cd "{wd}" && zip -q -X -r "{os.path.abspath(out)}" .', shell=True, check=True)
    shutil.rmtree(wd)

# ------------------------------------------------------------------ mode 2 : structure (openpyxl, classeurs simples)
def guess_table(ws):
    """Repère la ligne d'en-tête : première ligne à ≥ 2 libellés texte suivie d'au moins une ligne remplie."""
    maxr, maxc = ws.max_row, ws.max_column
    for r in range(1, min(maxr, 30) + 1):
        vals = [ws.cell(r, c).value for c in range(1, maxc + 1)]
        txt = [v for v in vals if isinstance(v, str) and v.strip()]
        filled = [v for v in vals if v not in (None, '')]
        if len(txt) >= 2 and len(txt) >= 0.6 * len(filled) and r < maxr:
            nxt = [ws.cell(r + 1, c).value for c in range(1, maxc + 1)]
            if sum(v not in (None, '') for v in nxt) >= 1:
                cols = [c for c in range(1, maxc + 1) if ws.cell(r, c).value not in (None, '')]
                last = r
                for rr in range(r + 1, maxr + 1):
                    if any(ws.cell(rr, c).value not in (None, '') for c in range(cols[0], cols[-1] + 1)): last = rr
                return dict(header_row=r, first_col=cols[0], last_col=cols[-1], last_row=last)
    return None

def structure(src, out, params):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
    wb = openpyxl.load_workbook(src)
    fill = lambda h: PatternFill('solid', start_color='FF' + h, end_color='FF' + h)
    thin = Side(style='thin', color='FF' + LINE); dot = Side(style='dotted', color='FF7F7F7F'); white = Side(style='thin', color='FFFFFFFF')
    for ws in wb.worksheets:
        pr = params.get('sheets', {}).get(ws.title, {})
        if pr.get('skip'): continue
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = 'FF' + GREEN
        # 1. palette + police sur toutes les cellules existantes
        for row in ws.iter_rows():
            for c in row:
                f = c.font
                col = f.color.rgb if f.color is not None and isinstance(f.color.rgb, str) else None
                newc = ('FF' + classify(col[-6:], 'text')) if col else None
                c.font = Font(name='Arial', size=f.size, bold=f.bold, italic=f.italic, underline=f.underline, strike=f.strike, vertAlign=f.vertAlign,
                              color=newc if newc else (f.color if f.color is not None and f.color.rgb is None else None))
                if c.fill is not None and c.fill.fill_type == 'solid' and isinstance(c.fill.fgColor.rgb, str) and c.fill.fgColor.rgb not in ('00000000',):
                    c.fill = fill(classify(c.fill.fgColor.rgb[-6:], 'fill'))
        t = {**(guess_table(ws) or {}), **{k: v for k, v in pr.items() if k in ('header_row', 'first_col', 'last_col', 'last_row')}}
        if not t.get('header_row'): continue
        hr, c0, c1, lr = t['header_row'], t['first_col'], t['last_col'], t['last_row']
        # 2. en-tête vert, texte blanc gras centré (Q&A Dune)
        for c in range(c0, c1 + 1):
            cell = ws.cell(hr, c)
            cell.fill = fill(GREEN)
            cell.font = Font(name='Arial', size=cell.font.size or 10, bold=True, color='FFFFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = Border(left=white, right=white, top=thin, bottom=thin)
        # 3. lignes : séparateurs pointillés, colonne d'index grise, lignes de total surlignées
        idx = pr.get('index_col', all(isinstance(ws.cell(r, c0).value, (int, float)) or ws.cell(r, c0).value is None for r in range(hr + 1, lr + 1)))
        for r in range(hr + 1, lr + 1):
            label = ' '.join(str(ws.cell(r, c).value or '') for c in range(c0, min(c1, c0 + 2) + 1)).lower()
            total = bool(re.search(r'\b(total|sous-total|ensemble|synth[eè]se)\b', label)) or r in pr.get('total_rows', [])
            for c in range(c0, c1 + 1):
                cell = ws.cell(r, c)
                cell.border = Border(bottom=dot, top=Side(style='thin', color='FF' + GREEN) if total else None)
                if cell.alignment is None or not cell.alignment.vertical:
                    cell.alignment = Alignment(horizontal=cell.alignment.horizontal, vertical='center', wrap_text=cell.alignment.wrap_text)
                if total:
                    cell.fill = fill(SOFT); cell.font = Font(name='Arial', size=cell.font.size, bold=True, italic=cell.font.italic, color='FF' + DEEP)
                elif idx and c == c0:
                    cell.fill = fill('D9D9D9'); cell.font = Font(name='Arial', size=cell.font.size, bold=True, color='FF' + DEEP)
                    cell.alignment = Alignment(horizontal='center', vertical='center')
        # 4. bloc titre si des lignes libres existent au-dessus du tableau (jamais d'insertion de ligne : formules protégées)
        title = pr.get('title')
        if title and hr >= 4 and all(ws.cell(r, c).value in (None, '') for r in range(1, hr - 1) for c in range(1, c1 + 1)):
            tr = hr - 3 if hr >= 5 else 1
            span = max(c0, min(c1, c0 + 5))
            for r in (tr, tr + 1):
                for c in range(c0, span + 1): ws.cell(r, c).fill = fill(GREEN)
            ws.cell(tr, c0).value = title; ws.cell(tr, c0).font = Font(name='Arial', size=15, color='FFFFFFFF')
            if pr.get('subtitle'):
                ws.cell(tr + 1, c0).value = pr['subtitle']; ws.cell(tr + 1, c0).font = Font(name='Arial', size=11, color='FF' + 'D6D2CB')
            ws.row_dimensions[tr].height = 24; ws.row_dimensions[tr + 1].height = 18
        # 5. confort : figer sous l'en-tête, filtre, marge colonne A
        if pr.get('freeze', True) and ws.freeze_panes is None:
            ws.freeze_panes = ws.cell(hr + 1, c0)
        if pr.get('autofilter') and not ws.auto_filter.ref:
            ws.auto_filter.ref = f'{get_column_letter(c0)}{hr}:{get_column_letter(c1)}{lr}'
        if c0 > 1 and all(ws.cell(r, 1).value in (None, '') for r in range(1, ws.max_row + 1)):
            ws.column_dimensions['A'].width = 2.4
        ws.oddHeader.right.text = params.get('print_header', ''); ws.oddHeader.right.font = 'Arial,Bold'; ws.oddHeader.right.color = GREEN
        ws.oddFooter.left.text = params.get('print_footer', ''); ws.oddFooter.left.font = 'Arial'; ws.oddFooter.left.color = SUB
        ws.oddFooter.right.text = '&P / &N'; ws.oddFooter.right.font = 'Arial,Bold'; ws.oddFooter.right.color = GREEN
    wb.save(out)

# ------------------------------------------------------------------ inspection / contrôle
def inspect(path):
    import openpyxl
    r = risky_parts(path)
    print('Éléments sensibles (graphiques, TCD, images, macros) :', r or 'aucun')
    print('=> mode conseillé :', 'palette (XML) uniquement' if r else 'structure + palette')
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        nf = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('='))
        print(f'\n== {ws.title} | plage {ws.dimensions} | formules {nf} | fusions {len(ws.merged_cells.ranges)} | volet figé {ws.freeze_panes}')
        print('   tableau détecté :', guess_table(ws))
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10), values_only=True):
            print('   ', [str(v)[:18] for v in row[:10]])

def check(a, b):
    import openpyxl
    A, B = openpyxl.load_workbook(a), openpyxl.load_workbook(b)
    diffs = 0
    for ws in A.worksheets:
        wb2 = B[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if c.value != wb2[c.coordinate].value:
                    diffs += 1
                    if diffs < 15: print('Différence', ws.title, c.coordinate, repr(c.value)[:40], '->', repr(wb2[c.coordinate].value)[:40])
    print('Cellules modifiées :', diffs, '(seul le bloc titre ajouté peut en créer)')

def apply(src, out, params):
    r = risky_parts(src)
    if r and not params.get('force_structure'):
        palette_xml(src, out); print('Mode palette (XML) : graphiques / TCD / images préservés ->', out); return
    tmp = out + '.tmp.xlsx'
    structure(src, tmp, params); palette_xml(tmp, out); os.remove(tmp)
    print('Mode structure + palette ->', out)

if __name__ == '__main__':
    cmd = sys.argv[1]
    if cmd == 'inspect': inspect(sys.argv[2])
    elif cmd == 'apply': apply(sys.argv[2], sys.argv[3], json.load(open(sys.argv[4], encoding='utf-8')) if len(sys.argv) > 4 else {})
    elif cmd == 'check': check(sys.argv[2], sys.argv[3])
