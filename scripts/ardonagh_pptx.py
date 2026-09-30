#!/usr/bin/env python3
"""Charte Ardonagh pour tout PowerPoint (modèle : one-pager M&A du groupe). Le fond n'est jamais modifié.
Usage :
  python3 ardonagh_pptx.py inspect in.pptx
  python3 ardonagh_pptx.py apply in.pptx out.pptx params.json
  python3 ardonagh_pptx.py check in.pptx out.pptx
"""
import re, os, sys, json, copy, zipfile, shutil, tempfile, subprocess, hashlib
from collections import Counter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from ardonagh_word import classify, recolor_png, GREEN, DEEP, MID, SOFT, LINE, SUB, GOLD, SAND, RED, BAR, THEME
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
IN = 914400
REF_W, REF_H = 10 * IN, 5.625 * IN          # géométrie du modèle Ardonagh (16:9, 10 x 5,625 pouces)
ADDED = 'ARD_'                               # préfixe des formes ajoutées par le moteur
SCALE = [1.0]                                # rapport largeur diapo / largeur modèle, pour les tailles de police

# ------------------------------------------------------------------ utilitaires
def rgb(h): return RGBColor.from_string(h)

def rect(shapes, name, x, y, w, h, fill, line=None):
    s = shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x), int(y), int(w), int(h))
    s.name = ADDED + name
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(1)
    else: s.line.fill.background()
    s.shadow.inherit = False
    return s

def hline(shapes, name, x, y, w, color='FFFFFF', width=1):
    c = shapes.add_connector(1, int(x), int(y), int(x + w), int(y))
    c.name = ADDED + name; c.line.color.rgb = rgb(color); c.line.width = Pt(width)
    return c

def textbox(shapes, name, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines = [(texte, taille_pt, gras, italique, couleur, capitales)]"""
    tb = shapes.add_textbox(int(x), int(y), int(w), int(h)); tb.name = ADDED + name
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    for i, (t, sz, b, it, col, caps) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = t
        f = r.font; f.name = 'Arial'; f.size = Pt(round(sz * SCALE[0], 1)); f.bold = b; f.italic = it; f.color.rgb = rgb(col)
        if caps: r._r.get_or_add_rPr().set('cap', 'all')
    return tb

def send_to_back(shape):
    tree = shape._element.getparent(); tree.remove(shape._element); tree.insert(2, shape._element)

def texts_by_slide(pres, include_added=False):
    res = []
    for s in pres.slides:
        c = Counter()
        for sp in s.shapes:
            if sp.name.startswith(ADDED) and not include_added: continue
            for t in sp._element.iter('{%s}t' % A):
                if t.text and t.text.strip(): c[t.text] += 1
        res.append(c)
    return res

def is_title_ph(sh):
    return sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)

# ------------------------------------------------------------------ inspection
def inspect(path):
    p = Presentation(path)
    print(f'Format : {p.slide_width / IN:.2f} x {p.slide_height / IN:.2f} pouces | {len(p.slides)} diapositives')
    for i, s in enumerate(p.slides, 1):
        title = next((sh.text_frame.text for sh in s.shapes if is_title_ph(sh) and sh.has_text_frame), '')
        kinds = Counter(str(sh.shape_type).split('.')[-1].split(' ')[0] for sh in s.shapes)
        top = min([sh.top for sh in s.shapes if sh.top is not None and not is_title_ph(sh)] or [0]) / IN
        print(f'{i:>2} | mise en page « {s.slide_layout.name} » | titre : {title[:60]!r} | formes : {dict(kinds)} | plus haute forme hors titre : {top:.2f} po')
    for m in p.slide_masters:
        pics = [sh.name for sh in m.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
        print('Masque :', m.name if hasattr(m, 'name') else '', '| images décoratives :', pics)

# ------------------------------------------------------------------ couverture et page de fin
def add_picture_safe(shapes, path, x, y, w=None, h=None, name='img'):
    if not path or not os.path.exists(path): return None
    pic = shapes.add_picture(path, int(x), int(y), int(w) if w else None, int(h) if h else None); pic.name = ADDED + name
    return pic

def build_cover(slide, c, W, H, assets, report):
    sx, sy = W / REF_W, H / REF_H
    # on garde les textes et un éventuel tableau de suivi ; images et formes décoratives sont retirées
    texts, tables = [], []
    for sh in list(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append((sh.top or 0, sh.left or 0, sh, [p.text for p in sh.text_frame.paragraphs if p.text.strip()]))
        elif sh.shape_type == MSO_SHAPE_TYPE.TABLE:
            tables.append(sh); continue
        elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            report.append(f'Couverture : image « {sh.name} » retirée (remplacée par le logo et la photo Ardonagh)')
        sh._element.getparent().remove(sh._element)
    texts.sort(key=lambda t: (t[0], t[1]))
    title_sh = next((t for t in texts if is_title_ph(t[2])), texts[0] if texts else None)
    kicker = c.get('kicker') or (title_sh[3][0] if title_sh else '')
    title = c.get('title') if 'title' in c else (' '.join(title_sh[3][1:]) if title_sh and len(title_sh[3]) > 1 else '')
    others = c.get('subs') if 'subs' in c else [l for t in texts if t is not title_sh for l in t[3] if l != c.get('date')]
    # fond
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = rgb(GREEN)
    sh = slide.shapes
    rect(sh, 'CoverGreen', 0, 0, W, H, GREEN)
    photo_w = 2.58 * IN * sx
    if add_picture_safe(sh, os.path.join(assets, 'photo.jpg'), 0, 0, photo_w, H, 'CoverPhoto') is not None:
        c_ = sh.add_connector(1, int(photo_w), 0, int(photo_w), int(H)); c_.name = ADDED + 'CoverLine'; c_.line.color.rgb = rgb('FFFFFF'); c_.line.width = Pt(0.75)
    else:
        photo_w = 0; report.append('Couverture sans photo (assets/photo.jpg absent)')
    zone_x, zone_w = photo_w + (4.13 * IN * sx - 2.58 * IN * sx), 4.30 * IN * sx
    if photo_w == 0: zone_x = (W - zone_w) / 2
    logo_w = 3.42 * IN * sx
    if add_picture_safe(sh, os.path.join(assets, 'logo_white.png'), zone_x + (zone_w - logo_w) / 2, 1.84 * IN * sy, logo_w, None, 'CoverLogo') is None:
        textbox(sh, 'CoverWordmark', zone_x, 1.9 * IN * sy, zone_w, 1.1 * IN * sy, [('THE', 9, False, False, 'FFFFFF', True), ('ARDONAGH', 30, False, False, 'FFFFFF', True), ('GROUP', 9, False, False, 'FFFFFF', True)], PP_ALIGN.CENTER)
        report.append('Couverture sans logo (assets/logo_white.png absent) : mot-marque typographique utilisé')
    y1, y2 = 3.80 * IN * sy, 4.81 * IN * sy
    hline(sh, 'CoverRuleTop', zone_x, y1, zone_w); hline(sh, 'CoverRuleBottom', zone_x, y2, zone_w)
    lines = [(c['strap'], 7, True, False, GOLD, True), (kicker, 12, True, False, 'FFFFFF', True)]
    if title: lines.append((title, 12, False, True, 'FFFFFF', False))
    lines += [(o, 8, False, False, SAND, False) for o in others]
    textbox(sh, 'CoverText', zone_x, y1 + 0.08 * IN * sy, zone_w, y2 - y1 - 0.12 * IN * sy, lines, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    if c.get('date'):
        textbox(sh, 'CoverDate', zone_x, y2 + 0.06 * IN * sy, zone_w, 0.25 * IN * sy, [(c['date'], 8, True, False, 'FFFFFF', True)], PP_ALIGN.RIGHT)
    for t in tables:   # tableau de suivi en haut à droite, filets blancs (façon « Template Version Control »)
        t.left, t.top = int(W - t.width - 0.26 * IN * sx), int(0.22 * IN * sy)
        for row in t.table.rows:
            for cell in row.cells:
                cell.fill.background()
                for p in cell.text_frame.paragraphs:
                    for r in p.runs: r.font.color.rgb = rgb('FFFFFF'); r.font.name = 'Arial'
        tp = t._element.find('.//{%s}tblPr' % A)
        if tp is not None:
            for k in ('firstRow', 'bandRow'): tp.attrib.pop(k, None)
            sid = tp.find('{%s}tableStyleId' % A)
            if sid is not None: tp.remove(sid)
        for tc in t._element.iter('{%s}tc' % A):
            tcpr = tc.find('{%s}tcPr' % A)
            if tcpr is None: tcpr = etree.SubElement(tc, '{%s}tcPr' % A)
            for tag in ('lnL', 'lnR', 'lnT', 'lnB'):
                for e in tcpr.findall('{%s}%s' % (A, tag)): tcpr.remove(e)
            for i, tag in enumerate(('lnL', 'lnR', 'lnT', 'lnB')):
                ln = etree.fromstring(f'<a:{tag} xmlns:a="{A}" w="6350"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:{tag}>')
                tcpr.insert(i, ln)
        sp = t._element; sp.getparent().remove(sp); slide.shapes._spTree.append(sp)

def blank_layout(pres):
    return min(pres.slide_layouts, key=lambda l: len(l.placeholders))

def add_closing(pres, W, H, assets):
    s = pres.slides.add_slide(blank_layout(pres))
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = rgb(GREEN)
    rect(s.shapes, 'BackGreen', 0, 0, W, H, GREEN)
    lw = 3.15 * IN * W / REF_W
    if add_picture_safe(s.shapes, os.path.join(assets, 'logo_white.png'), (W - lw) / 2, 2.25 * IN * H / REF_H, lw, None, 'BackLogo') is None:
        textbox(s.shapes, 'BackWordmark', 0, H * 0.4, W, H * 0.2, [('THE ARDONAGH GROUP', 20, False, False, 'FFFFFF', True)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------------------ pages de contenu
def build_content(slide, idx, c, W, H, report):
    sx, sy = W / REF_W, H / REF_H
    band_h, strip_h = 0.51 * IN * sy, 0.09 * IN * sy
    bar_y, bar_h = H - 0.09 * IN * sy, 0.09 * IN * sy
    sh = slide.shapes
    # pieds de page d'origine : texte récupéré dans le pied Ardonagh, placeholders retirés
    kept_footer = []
    for s in list(sh):
        if s.is_placeholder and s.placeholder_format.type in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE):
            if s.placeholder_format.type == PP_PLACEHOLDER.FOOTER and s.text_frame.text.strip(): kept_footer.append(s.text_frame.text.strip())
            s._element.getparent().remove(s._element)
    # contenu qui empiéterait sur le bandeau : décalé vers le bas s'il reste de la place
    limit = band_h + 0.12 * IN * sy
    for s in sh:
        if s.name.startswith(ADDED) or is_title_ph(s) or s.top is None or s.height is None: continue
        if s.top < limit:
            d = limit - s.top
            if s.top + s.height + d <= bar_y - 0.2 * IN * sy: s.top = int(s.top + d)
            else: report.append(f'Diapo {idx} : « {s.name} » chevauche le bandeau et ne peut pas descendre sans sortir de la page')
    band = rect(sh, 'Band', 0, 0, W, band_h, GREEN); strip = rect(sh, 'BandStrip', 0, 0, W, strip_h, DEEP)
    bar = rect(sh, 'BottomBar', 0, bar_y, W, bar_h, BAR)
    for s in (bar, strip, band): send_to_back(s)
    # titre dans le bandeau : blanc, capitales, Arial 15
    title = next((s for s in sh if is_title_ph(s)), None)
    if title is not None and title.has_text_frame and title.text_frame.text.strip():
        title.left, title.top, title.width, title.height = int(0.39 * IN * sx), int(strip_h), int(W - 0.78 * IN * sx), int(band_h - strip_h)
        tf = title.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = tf._txBody.find('{%s}bodyPr' % A)
        for e in list(bp): bp.remove(e)
        for k in ('lIns', 'tIns', 'rIns', 'bIns'): bp.set(k, '0')
        n = len(tf.text)
        size = 15 if n <= 60 else 12 if n <= 90 else 10
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for r in p.runs:
                r.font.size = Pt(round(size * SCALE[0], 1)); r.font.bold = False; r.font.name = 'Arial'; r.font.color.rgb = rgb('FFFFFF')
                r._r.get_or_add_rPr().set('cap', 'all')
        tree = title._element.getparent(); tree.remove(title._element); tree.insert(list(tree).index(bar._element) + 1, title._element)
    elif c.get('band'):
        textbox(sh, 'BandText', 0.39 * IN * sx, strip_h, W - 0.78 * IN * sx, band_h - strip_h, [(c['band'], 15, False, False, 'FFFFFF', True)], PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    # pied : mention de groupe à gauche, numéro à droite
    foot = ' · '.join([c['footer']] + kept_footer)
    textbox(sh, 'Footer', 0.39 * IN * sx, bar_y - 0.24 * IN * sy, W - 1.3 * IN * sx, 0.16 * IN * sy, [(foot, 6.5, False, False, SUB, False)], PP_ALIGN.LEFT, MSO_ANCHOR.BOTTOM)
    num = textbox(sh, 'SlideNum', W - 0.75 * IN * sx, bar_y - 0.26 * IN * sy, 0.36 * IN * sx, 0.18 * IN * sy, [('#', 7.5, False, False, DEEP, False)], PP_ALIGN.RIGHT, MSO_ANCHOR.BOTTOM)
    r = num.text_frame.paragraphs[0].runs[0]._r
    fld = etree.fromstring(f'<a:fld xmlns:a="{A}" id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" type="slidenum"/>')
    fld.append(copy.deepcopy(r.find('{%s}rPr' % A))); t = etree.SubElement(fld, '{%s}t' % A); t.text = str(idx)
    r.addprevious(fld); r.getparent().remove(r)

# ------------------------------------------------------------------ palette (XML, sans perte)
def color_context(el):
    for a in el.iterancestors():
        n = etree.QName(a).localname
        if n in ('rPr', 'defRPr', 'endParaRPr', 'buClr'): return 'text'
        if n in ('ln', 'lnL', 'lnR', 'lnT', 'lnB'): return 'border'
        if n in ('spPr', 'tcPr', 'bg', 'bgPr', 'grpSpPr'): return 'fill'
    return 'fill'

def palette_part(xml_bytes):
    root = etree.fromstring(xml_bytes)
    for c in root.iter('{%s}srgbClr' % A):
        if any(etree.QName(a).localname in ('clrScheme', 'outerShdw', 'innerShdw') for a in c.iterancestors()): continue
        v = c.get('val', '')
        if re.fullmatch(r'[0-9A-Fa-f]{6}', v): c.set('val', classify(v, color_context(c)))
    for tag in ('latin', 'ea', 'cs'):
        for f in root.iter('{%s}%s' % (A, tag)):
            tf = f.get('typeface', '')
            if tf and not tf.startswith('+') and not re.match(r'(Wingdings|Symbol|Webdings|Marlett)', tf): f.set('typeface', 'Arial')
    return etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)

def palette_zip(src, out, skip_media_hashes, report, recolor=True, keep=()):
    wd = tempfile.mkdtemp(prefix='ardp_')
    with zipfile.ZipFile(src) as z:
        for n in z.namelist():
            if '..' in n or n.startswith('/'): continue
            z.extract(n, wd)
    for d, _, fs in os.walk(os.path.join(wd, 'ppt')):
        for f in fs:
            p = os.path.join(d, f); rel = os.path.relpath(p, wd).replace(os.sep, '/')
            if rel.startswith('ppt/theme/') and f.endswith('.xml'):
                t = open(p, encoding='utf-8').read()
                for k, v in {**THEME, 'dk1': '000000', 'lt1': 'FFFFFF'}.items():
                    t = re.sub(rf'(<a:{k}>)(.*?)(</a:{k}>)', rf'\1<a:srgbClr val="{v}"/>\3', t, flags=re.S)
                t = re.sub(r'(<a:(?:major|minor)Font><a:latin typeface=")[^"]*"', r'\1Arial"', t)
                open(p, 'w', encoding='utf-8').write(t)
            elif re.match(r'ppt/(slides|slideLayouts|slideMasters|charts|notesMasters|diagrams)/[^/]+\.xml$', rel):
                data = palette_part(open(p, 'rb').read())
                open(p, 'wb').write(data)
            elif recolor and rel.startswith('ppt/media/') and f.lower().endswith('.png') and f not in keep:
                if hashlib.sha1(open(p, 'rb').read()).hexdigest() in skip_media_hashes: continue
                if recolor_png(p): report.append(f'Schéma recoloré : {f}')
    if os.path.exists(out): os.remove(out)
    subprocess.run(f'cd "{wd}" && zip -q -X -r "{os.path.abspath(out)}" .', shell=True, check=True)
    shutil.rmtree(wd)

# ------------------------------------------------------------------ application
def apply(src, out, c):
    report = []
    assets = c.get('assets_dir', '')
    pres = Presentation(src); W, H = pres.slide_width, pres.slide_height
    SCALE[0] = max(0.8, min(W / REF_W, H / REF_H * 1.0))
    n = len(pres.slides)
    cover = int(c.get('cover_slide', 1)); skip = set(c.get('skip_slides', []))
    # images décoratives des masques et mises en page (ancien logo, bandeaux) retirées
    if c.get('clean_masters', True):
        for m in pres.slide_masters:
            for holder in [m] + list(m.slide_layouts):
                for s in list(holder.shapes):
                    if s.shape_type == MSO_SHAPE_TYPE.PICTURE and not s.is_placeholder:
                        report.append(f'Image décorative retirée du masque « {getattr(holder, "name", "")} » : {s.name}')
                        s._element.getparent().remove(s._element)
    for i, s in enumerate(pres.slides, 1):
        if i in skip: continue
        if i == cover and c.get('cover'): build_cover(s, c['cover'], W, H, assets, report)
        else: build_content(s, i, c, W, H, report)
    if c.get('closing', True): add_closing(pres, W, H, assets)
    tmp = out + '.tmp.pptx'; pres.save(tmp)
    hashes = {hashlib.sha1(open(os.path.join(assets, f), 'rb').read()).hexdigest() for f in ('logo_white.png', 'photo.jpg') if os.path.exists(os.path.join(assets, f))}
    palette_zip(tmp, out, hashes, report, c.get('recolor_images', True), set(c.get('keep_images', [])))
    os.remove(tmp)
    for r in report: print('-', r)
    print('OK ->', out)

def check(a, b):
    A_, B_ = texts_by_slide(Presentation(a)), texts_by_slide(Presentation(b), include_added=True)
    bad = 0
    for i, ca in enumerate(A_):
        cb = B_[i] if i < len(B_) else Counter()
        missing = ca - cb
        # sur la couverture, les textes sont redistribués : on compare le texte concaténé
        if missing:
            joined = ' '.join(k for k in cb.elements())
            missing = Counter({k: v for k, v in missing.items() if k not in joined})
        if missing:
            bad += 1; print(f'Diapo {i + 1} : textes absents -> {list(missing)[:5]}')
    print('Textes d\'origine tous présents' if not bad else f'{bad} diapositive(s) à corriger')

if __name__ == '__main__':
    cmd = sys.argv[1]
    if cmd == 'inspect': inspect(sys.argv[2])
    elif cmd == 'apply': apply(sys.argv[2], sys.argv[3], json.load(open(sys.argv[4], encoding='utf-8')))
    elif cmd == 'check': check(sys.argv[2], sys.argv[3])
